#!/usr/bin/env python3
"""
diag_excel.py — Diagnóstico de las 2 dudas de David (NO es el arnés, es la
sonda que informa de las cifras para luego fijar los tests):

  1. ¿El Excel v1 (datos oficiales del crudo) sigue intacto en el v2?
     -> para CADA celda con valor: ¿viene del crudo (mismo valor) o de
        mejorado (con cita verificada)?; y ¿alguna celda determinista
        CAMBIÓ respecto al crudo?
  2. ¿Hasta qué punto está normalizando el v2 los títulos entre años?
     -> % con titulo_normalizado, patrones de formato, topónimos pendientes.
"""
import json, re, sys
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[2]
CRUDO = RAIZ / "database" / "data" / "crudo"
MEJ = RAIZ / "database" / "data" / "mejorado"
XLSX = RAIZ / "database" / "data" / "ciaf_base_global.xlsx"
GUIA = RAIZ / "database" / "data" / "guia_campos.json"


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


# ------------------------------------------------------------------ datos
guia = json.loads(GUIA.read_text(encoding="utf-8"))
refs = [str(g.get("ref")) for g in guia
        if str(g.get("ref", "")).strip()[:1].isdigit() and (g.get("campo") or "").strip()]

crudo = {}
for f in sorted(CRUDO.glob("*.json")):
    d = json.loads(f.read_text(encoding="utf-8"))
    crudo[Path(d.get("md") or f).stem] = d

mej = {}
for f in sorted(MEJ.glob("*.json")):
    d = json.loads(f.read_text(encoding="utf-8"))
    mej[d.get("stem") or f.stem] = d

wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["Informes"]
filas = list(ws.iter_rows(values_only=True))
cab = [str(c or "") for c in filas[0]]
i_md = cab.index("md") if "md" in cab else None
i_tit = cab.index("titulo") if "titulo" in cab else None
i_titn = cab.index("titulo_normalizado") if "titulo_normalizado" in cab else None
col = {r: cab.index(r) for r in refs if r in cab}
datos = [f for f in filas[1:] if any(f)]
wb.close()

# ---------------------------------------- 1) ¿v1 intacto dentro del v2?
celdas_total = 0
celdas_crudo = 0            # valor == valor del crudo  -> v1 intacto
celdas_mejor = 0            # valor justificado por mejorado verificado
celdas_sin_fuente = []       # ¡¡valor en el Excel que nadie respalda!!
celdas_deterministas_modificadas = []  # crudo tenia valor y el Excel OTRO
huecos_rellenados = 0       # crudo vacio y el Excel con valor (buena señal)
mej_no_verificado = 0

for f in datos:
    stem = Path(f[i_md]).stem if i_md and f[i_md] else ""
    dc = crudo.get(stem, {})
    campos = dc.get("campos", {})
    dm = mej.get(stem, {}).get("campos", {})
    for r, i in col.items():
        v = norm(f[i])
        if not v:
            continue
        celdas_total += 1
        vc = norm((campos.get(r) or {}).get("valor_fuente"))
        if vc:
            if vc == v:
                celdas_crudo += 1
            else:
                celdas_deterministas_modificadas.append((stem, r, vc[:60], v[:60]))
            continue
        # crudo vacio -> el Excel tiene algo
        m = dm.get(r) or {}
        if norm(m.get("valor")) == v and m.get("verificado"):
            celdas_mejor += 1
            huecos_rellenados += 1
        else:
            if m:
                mej_no_verificado += 1
            celdas_sin_fuente.append((stem, r, v[:70]))

print("=== 1) INTEGRIDAD v1 dentro del v2 ===")
print(f"celdas con valor en el Excel      : {celdas_total}")
print(f"  == valor del crudo (v1 intacto) : {celdas_crudo}"
      f" ({100*celdas_crudo/max(1,celdas_total):.1f}%)")
print(f"  rellenadas por mejorado verif.  : {celdas_mejor}")
print(f"  !! SIN FUENTE (nadie las apoya) : {len(celdas_sin_fuente)}")
print(f"  !! determinista MODIFICADA      : {len(celdas_deterministas_modificadas)}")
print(f"  mejorado con valor no verificado: {mej_no_verificado}")
for x in celdas_sin_fuente[:8]:
    print("     sin-fuente:", x)
for x in celdas_deterministas_modificadas[:8]:
    print("     modificada:", x)

# ------------------------------------------------------ 2) normalización
print()
print("=== 2) NORMALIZACIÓN DE TÍTULOS (v2) ===")
tots = 0
con_norm = 0
iguales = 0
patrones = Counter()
pendientes = []
for f in datos:
    t = norm(f[i_tit]) if i_tit is not None else ""
    tn = norm(f[i_titn]) if i_titn is not None else ""
    if not t:
        continue
    tots += 1
    if tn:
        con_norm += 1
        if tn == t:
            iguales += 1
    # ¿patrón de formato del CRUDO? (para medir la variación entre años)
    if re.match(r"^[0-9N][^\n]{0,12}n[ºo]\s*\d+/\d{4}", t, re.I):
        patrones["expediente-primero"] += 1
    elif re.match(r"^investigaci[oó]n", t, re.I):
        patrones["investigación..."] += 1
    elif re.match(r"^(informe|if|ciaf)\b", t, re.I):
        patrones["informe/if/ciaf..."] += 1
    elif t.isupper():
        patrones["TODO MAYÚSCULAS"] += 1
    else:
        patrones["otro"] += 1
    # topónimos en minúscula tras "en ... de" (medir, no arreglar a ciegas)
    if re.search(r"\b(en|desde|hacia|cerca de)\s+[a-záéíóúñ][a-záéíóúñ]+", t):
        pendientes.append((norm(f[i_md]) if i_md and f[i_md] else "", t[:95]))

print(f"títulos                     : {tots}")
print(f"con titulo_normalizado      : {con_norm} ({100*con_norm/max(1,tots):.1f}%)")
print(f"  de ellos, IGUALES al crudo: {iguales} ({100*iguales/max(1,con_norm):.1f}%)"
      "  <- si es ~100% el v2 NO está tocando nada")
print("patrones de formato en el crudo:")
for p, n in patrones.most_common():
    print(f"  {p:24} {n:4} ({100*n/max(1,tots):.1f}%)")
print(f"títulos con posible topónimo en minúscula: {len(pendientes)}")
for x in pendientes[:8]:
    print("     ", x)

# --------------------------------------------- 3) huecos que quedan
print()
print("=== 3) HUECOS QUE QUEDAN POR CAMPO (ref con <100%) ===")
vacias = Counter()
llenar = Counter()
for f in datos:
    stem = Path(f[i_md]).stem if i_md and f[i_md] else ""
    dm = mej.get(stem, {}).get("campos", {})
    for r, i in col.items():
        if not norm(f[i]):
            vacias[r] += 1
            if dm.get(r):
                llenar[r] += 1
tot_col = len(col)
print(f"columnas de guía: {tot_col} · filas: {len(datos)} · celdas vacías: {sum(vacias.values())}")
top = vacias.most_common(8)
for r, n in top:
    campo = next((g.get("campo") for g in guia if str(g.get("ref")) == r), "")
    print(f"  {r:8} vacías {n:4}/{len(datos)}  {str(campo)[:52]}")
